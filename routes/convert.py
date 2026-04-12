# -*- coding: utf-8 -*-
import os
import io
import json
import base64
import requests
from flask import Blueprint, request, jsonify, send_file, render_template
from werkzeug.utils import secure_filename

convert_bp = Blueprint('convert', __name__)

# Configuration
UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'uploads')
CONVERT_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'tmp', 'convert')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(CONVERT_FOLDER, exist_ok=True)

ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'pdf'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

@convert_bp.route('/convert')
def index():
    """Render convert page"""
    return render_template('convert.html')

@convert_bp.route('/convert/process', methods=['POST'])
def process():
    """Handle file conversion"""
    if 'file' not in request.files:
        return jsonify({'error': 'Không tìm thấy file tải lên'}), 400
    
    file = request.files['file']
    convert_type = request.form.get('type', '')
    
    if file.filename == '':
        return jsonify({'error': 'Chưa chọn file'}), 400
    
    if not allowed_file(file.filename):
        return jsonify({'error': 'Định dạng file không được hỗ trợ. Chỉ chấp nhận: JPG, PNG, PDF'}), 400
    
    filename = secure_filename(file.filename)
    filepath = os.path.join(UPLOAD_FOLDER, filename)
    file.save(filepath)
    
    try:
        if convert_type == 'img_excel':
            return convert_image_to_excel_ocr(filepath)
        elif convert_type == 'pdf_excel':
            return convert_image_to_excel_ocr(filepath)  # Dùng chung
        elif convert_type == 'pdf_word':
            return convert_image_to_word_ocr(filepath)
        elif convert_type == 'img_word':
            return convert_image_to_word_ocr(filepath)
        elif convert_type == 'img_pdf':
            return convert_image_to_pdf_simple(filepath)
        else:
            return jsonify({'error': 'Loại chuyển đổi không hợp lệ'}), 400
    except Exception as e:
        print(f"Conversion error: {e}")
        return jsonify({'error': f'Lỗi xử lý: {str(e)}'}), 500
    finally:
        if os.path.exists(filepath):
            try:
                os.remove(filepath)
            except:
                pass

def convert_image_to_excel_ocr(filepath):
    """Convert image to Excel using OCR.space API (free)"""
    try:
        from openpyxl import Workbook
        
        # Gửi file đến OCR.space API (free tier)
        with open(filepath, 'rb') as f:
            files = {'file': f}
            data = {'language': 'vie', 'isTable': 'true'}
            response = requests.post(
                'https://api.ocr.space/parse/image',
                files=files,
                data=data,
                timeout=30
            )
        
        result = response.json()
        
        if result.get('IsErroredOnProcessing'):
            return jsonify({'error': result.get('ErrorMessage', ['Lỗi xử lý'])[0]}), 400
        
        # Lấy kết quả text
        text_results = result.get('ParsedResults', [])
        if not text_results:
            return jsonify({'error': 'Không nhận dạng được văn bản'}), 400
        
        text = text_results[0].get('ParsedText', '')
        
        # Tạo Excel với dữ liệu text
        wb = Workbook()
        ws = wb.active
        ws.title = "Dữ liệu OCR"
        
        # Ghi text vào cells (mỗi dòng là một dòng text)
        lines = text.split('\n')
        for row_idx, line in enumerate(lines, 1):
            if line.strip():
                # Chia theo tab hoặc space nhiều
                cells = line.split('\t')
                if len(cells) == 1:
                    cells = line.split()  # Chia theo space
                for col_idx, cell in enumerate(cells, 1):
                    ws.cell(row=row_idx, column=col_idx, value=cell.strip())
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='bang_du_lieu.xlsx'
        )
        
    except Exception as e:
        print(f"OCR error: {e}")
        return jsonify({'error': f'Lỗi OCR: {str(e)}'}), 500

def convert_image_to_word_ocr(filepath):
    """Convert image to Word using OCR.space API"""
    try:
        from docx import Document
        
        # Gửi file đến OCR.space API
        with open(filepath, 'rb') as f:
            files = {'file': f}
            data = {'language': 'vie'}
            response = requests.post(
                'https://api.ocr.space/parse/image',
                files=files,
                data=data,
                timeout=30
            )
        
        result = response.json()
        
        if result.get('IsErroredOnProcessing'):
            return jsonify({'error': result.get('ErrorMessage', ['Lỗi xử lý'])[0]}), 400
        
        text_results = result.get('ParsedResults', [])
        if not text_results:
            return jsonify({'error': 'Không nhận dạng được văn bản'}), 400
        
        text = text_results[0].get('ParsedText', '')
        
        # Tạo Word document
        doc = Document()
        doc.add_heading('Văn bản nhận dạng từ ảnh', 0)
        
        lines = text.split('\n')
        for line in lines:
            if line.strip():
                doc.add_paragraph(line.strip())
        
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            as_attachment=True,
            download_name='van_ban.docx'
        )
        
    except Exception as e:
        print(f"OCR error: {e}")
        return jsonify({'error': f'Lỗi OCR: {str(e)}'}), 500

def convert_image_to_pdf_simple(filepath):
    """Convert image to PDF using PIL"""
    try:
        from PIL import Image
        
        img = Image.open(filepath)
        
        if img.mode in ('RGBA', 'LA', 'P'):
            background = Image.new('RGB', img.size, (255, 255, 255))
            if img.mode == 'P':
                img = img.convert('RGBA')
            if img.mode in ('RGBA', 'LA'):
                background.paste(img, mask=img.split()[-1])
                img = background
            else:
                img = img.convert('RGB')
        elif img.mode != 'RGB':
            img = img.convert('RGB')
        
        output = io.BytesIO()
        img.save(output, format='PDF')
        output.seek(0)
        
        return send_file(
            output,
            mimetype='application/pdf',
            as_attachment=True,
            download_name='chuyen_doi.pdf'
        )
        
    except Exception as e:
        print(f"Image to PDF error: {e}")
        return jsonify({'error': f'Lỗi chuyển đổi ảnh sang PDF: {str(e)}'}), 500
